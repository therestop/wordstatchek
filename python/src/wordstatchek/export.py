from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.table import Table, TableStyleInfo

from .models import CheckResult

HEADERS = ["phrase", "total_count", "status", "error", "checked_at"]


def export_results(results: dict[str, CheckResult], output: Path, *, xlsx: bool = True) -> None:
    output.mkdir(parents=True, exist_ok=True)
    values = list(results.values())
    _write_csv(output / "wordstat_all.csv", values)
    _write_csv(output / "wordstat_nonzero.csv", [x for x in values if x.status == "nonzero"])
    _write_txt(output / "wordstat_nonzero.txt", [x.phrase for x in values if x.status == "nonzero"])
    _write_txt(output / "wordstat_zero.txt", [x.phrase for x in values if x.status == "zero"])
    _write_txt(
        output / "wordstat_errors.txt",
        [f"{x.phrase}\t{x.error}" for x in values if x.status == "error"],
    )
    if xlsx:
        _write_xlsx(output / "wordstat_results.xlsx", values)


def _write_csv(path: Path, values: list[CheckResult]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=HEADERS)
        writer.writeheader()
        writer.writerows(item.to_dict() for item in values)


def _write_txt(path: Path, lines: list[str]) -> None:
    path.write_text("\n".join(lines) + ("\n" if lines else ""), encoding="utf-8")


def _write_xlsx(path: Path, values: list[CheckResult]) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    groups = {
        "Все": values,
        "Ненулевые": [x for x in values if x.status == "nonzero"],
        "Нулевые": [x for x in values if x.status == "zero"],
        "Ошибки": [x for x in values if x.status == "error"],
    }
    for title, items in groups.items():
        sheet = workbook.create_sheet(title)
        sheet.append(HEADERS)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for item in items:
            sheet.append([item.phrase, item.total_count, item.status, item.error, item.checked_at])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.column_dimensions["A"].width = 52
        sheet.column_dimensions["B"].width = 16
        sheet.column_dimensions["C"].width = 14
        sheet.column_dimensions["D"].width = 55
        sheet.column_dimensions["E"].width = 28
        if items:
            table = Table(displayName=f"Wordstat{len(workbook.worksheets)}", ref=sheet.dimensions)
            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False
            )
            sheet.add_table(table)
    workbook.save(path)

